import os
import re
import shutil
import sys
import zipfile


MAX_PDF_PAGES = 100
MAX_PPTX_PAGES = 50
MAX_RENDER_PIXELS = 40_000_000


def open_pdf(path):
    import pymupdf

    document = pymupdf.open(path)
    if document.is_encrypted:
        document.close()
        raise ValueError("Password-protected PDFs are not supported.")
    if document.page_count < 1 or document.page_count > MAX_PDF_PAGES:
        document.close()
        raise ValueError("PDF must contain between 1 and 100 pages.")
    return document


def pdf_to_docx(source, destination):
    from pdf2docx import Converter

    document = open_pdf(source)
    try:
        if not any(page.get_text().strip() for page in document):
            raise ValueError("This PDF has no selectable text. Scanned PDFs need OCR, which is not enabled.")
        page_count = document.page_count
    finally:
        document.close()

    converter = Converter(source)
    try:
        converter.convert(destination, start=0, end=page_count)
    finally:
        converter.close()


def safe_cell(value):
    value = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", str(value or ""))[:32767]
    if value.startswith(("=", "+", "-", "@")):
        value = "'" + value
    return value


def pdf_to_xlsx(source, destination):
    import pdfplumber
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.remove(workbook.active)
    found_text = False
    with pdfplumber.open(source) as pdf:
        if not 1 <= len(pdf.pages) <= MAX_PDF_PAGES:
            raise ValueError("PDF must contain between 1 and 100 pages.")
        for index, page in enumerate(pdf.pages, start=1):
            sheet = workbook.create_sheet("Page {}".format(index))
            tables = page.extract_tables()
            if tables:
                for table in tables:
                    for row in table:
                        sheet.append([safe_cell(value) for value in (row or [])])
                        found_text = found_text or any(str(value or "").strip() for value in row or [])
                    sheet.append([])
            else:
                text = page.extract_text() or ""
                for line in text.splitlines():
                    sheet.append([safe_cell(line)])
                    found_text = found_text or bool(line.strip())
        if not found_text:
            raise ValueError("This PDF has no selectable text to extract. Scanned PDFs need OCR.")
    workbook.save(destination)


def pdf_to_pptx(source, destination):
    import pymupdf
    from pptx import Presentation
    from pptx.util import Inches

    document = open_pdf(source)
    try:
        if document.page_count > MAX_PPTX_PAGES:
            raise ValueError("PDF-to-PowerPoint supports up to 50 pages.")
        presentation = Presentation()
        presentation.slide_width = Inches(13.333)
        presentation.slide_height = Inches(7.5)
        blank_layout = presentation.slide_layouts[6]
        slide_width = presentation.slide_width
        slide_height = presentation.slide_height
        for page in document:
            pixmap = page.get_pixmap(matrix=pymupdf.Matrix(1.25, 1.25), alpha=False)
            if pixmap.width * pixmap.height > 40_000_000:
                raise ValueError("A PDF page is too large to render safely.")
            image_path = os.path.join(os.path.dirname(destination), "page-{}.png".format(page.number))
            pixmap.save(image_path)
            slide = presentation.slides.add_slide(blank_layout)
            ratio = min(slide_width / pixmap.width, slide_height / pixmap.height)
            width = int(pixmap.width * ratio)
            height = int(pixmap.height * ratio)
            slide.shapes.add_picture(
                image_path,
                int((slide_width - width) / 2),
                int((slide_height - height) / 2),
                width=width,
                height=height,
            )
            os.remove(image_path)
        presentation.save(destination)
    finally:
        document.close()


def pdf_to_images(source, destination, image_format):
    import pymupdf

    if image_format not in ("jpg", "png"):
        raise ValueError("Unsupported PDF image format.")

    document = open_pdf(source)
    try:
        with zipfile.ZipFile(destination, "w", compression=zipfile.ZIP_STORED) as archive:
            for index, page in enumerate(document, start=1):
                page_area = page.rect.width * page.rect.height
                if page_area <= 0:
                    raise ValueError("A PDF page has invalid dimensions.")
                scale = min(2.0, (MAX_RENDER_PIXELS / page_area) ** 0.5)
                pixmap = page.get_pixmap(matrix=pymupdf.Matrix(scale, scale), alpha=False)
                while pixmap.width * pixmap.height > MAX_RENDER_PIXELS:
                    scale *= (MAX_RENDER_PIXELS / (pixmap.width * pixmap.height)) ** 0.5 * 0.99
                    pixmap = page.get_pixmap(matrix=pymupdf.Matrix(scale, scale), alpha=False)
                if image_format == "jpg":
                    image = pixmap.tobytes("jpg", jpg_quality=85)
                else:
                    image = pixmap.tobytes("png")
                archive.writestr("page-{:03d}.{}".format(index, image_format), image)
    finally:
        document.close()


def compress_pdf(source, destination):
    import pymupdf

    document = open_pdf(source)
    optimized_path = os.path.join(os.path.dirname(source), "optimized.pdf")
    try:
        document.rewrite_images(
            dpi_threshold=300,
            dpi_target=150,
            quality=75,
            lossy=True,
            lossless=True,
            bitonal=False,
            color=True,
            gray=True,
        )
        document.save(
            optimized_path,
            garbage=4,
            deflate=True,
            deflate_images=True,
            deflate_fonts=True,
            use_objstms=True,
            compression_effort=100,
        )
    finally:
        document.close()
    try:
        if os.path.getsize(optimized_path) < os.path.getsize(source):
            os.replace(optimized_path, destination)
        else:
            shutil.copyfile(source, destination)
    finally:
        if os.path.isfile(optimized_path):
            os.remove(optimized_path)


def main():
    if len(sys.argv) != 4:
        raise ValueError("Invalid converter arguments.")
    operation, source, destination = sys.argv[1:]
    if operation == "pdf-to-docx":
        pdf_to_docx(source, destination)
    elif operation == "pdf-to-xlsx":
        pdf_to_xlsx(source, destination)
    elif operation == "pdf-to-pptx":
        pdf_to_pptx(source, destination)
    elif operation == "pdf-to-jpg":
        pdf_to_images(source, destination, "jpg")
    elif operation == "pdf-to-png":
        pdf_to_images(source, destination, "png")
    elif operation == "compress-pdf":
        compress_pdf(source, destination)
    else:
        raise ValueError("Unsupported document conversion.")
    if not os.path.isfile(destination) or os.path.getsize(destination) < 1:
        raise ValueError("The converter did not create an output file.")


if __name__ == "__main__":
    try:
        main()
    except Exception as error:
        print("{}: {}".format(type(error).__name__, error), file=sys.stderr)
        sys.exit(1)
